"""AI Resume ATS Checker - Streamlit + Google Gemini Flash."""

import io
import json
import os
import time
from typing import List, Optional

import streamlit as st
from docx import Document
from google import genai
from google.genai import errors, types
from pydantic import BaseModel, Field
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-flash-latest"  # alias to the newest Flash model
# Tried in order if the chosen model stays overloaded or is unavailable.
FALLBACK_MODELS = ["gemini-flash-lite-latest", "gemini-2.5-flash"]
RETRYABLE_CODES = {429, 500, 502, 503, 504}  # temporary errors worth retrying
MAX_RETRIES = 3  # attempts per model
RETRY_BASE_DELAY = 2  # seconds; doubles after each failed attempt
MAX_RESUME_CHARS = 30_000
MAX_FILE_MB = 5


# ----------------------------- Response schema -----------------------------
class SectionScore(BaseModel):
    name: str = Field(description="Section name, e.g. Contact Info, Skills")
    score: int = Field(description="Score from 0 to 100")
    feedback: str = Field(description="One or two sentences of feedback")


class Improvement(BaseModel):
    priority: str = Field(description="High, Medium or Low")
    issue: str = Field(description="What is wrong or missing")
    suggestion: str = Field(description="Specific fix, with an example if useful")


class ResumeAnalysis(BaseModel):
    ats_score: int = Field(description="Overall ATS score from 0 to 100")
    summary: str = Field(description="Two or three sentence overall assessment")
    strengths: List[str]
    section_scores: List[SectionScore]
    keywords_found: List[str]
    keywords_missing: List[str]
    formatting_issues: List[str]
    improvements: List[Improvement]


# ------------------------------ File parsing -------------------------------
def extract_text(uploaded_file) -> str:
    """Extract plain text from a PDF or DOCX upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")


# ------------------------------- Gemini call -------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description.strip()}\n\n"
        "Score the resume against this job description. Missing keywords must "
        "come from the job description."
        if job_description.strip()
        else "No job description was given. Evaluate general ATS-readiness and "
        "list commonly expected keywords for the candidate's apparent field."
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyze the resume below and return a structured evaluation.

Scoring guidance (be honest and calibrated, do not inflate):
- Keyword relevance and match: 30%
- Formatting and ATS parseability (clear headings, no tables/columns noise, standard sections): 20%
- Quantified achievements and impact: 20%
- Section completeness (contact, summary, experience, education, skills): 15%
- Grammar, clarity and conciseness: 15%

For section_scores include: Contact Info, Summary, Experience, Education, Skills, Formatting.
Give 5-8 improvements ordered by priority. Be specific and actionable.
Only use information present in the resume; never invent experience.

{jd_block}

RESUME:
\"\"\"
{resume_text[:MAX_RESUME_CHARS]}
\"\"\"
"""


def _generate(client, model: str, prompt: str) -> ResumeAnalysis:
    """One Gemini request, returned as a validated ResumeAnalysis."""
    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=ResumeAnalysis,
            temperature=0.2,
        ),
    )
    parsed: Optional[ResumeAnalysis] = getattr(response, "parsed", None)
    if isinstance(parsed, ResumeAnalysis):
        return parsed
    # Fallback: parse the raw JSON text ourselves.
    text = (response.text or "").strip()
    if text.startswith("```"):
        text = text.strip("`")
        text = text.split("\n", 1)[1] if "\n" in text else text
    return ResumeAnalysis.model_validate(json.loads(text))


def analyze_resume(
    api_key: str, model: str, resume_text: str, job_description: str
) -> ResumeAnalysis:
    """Analyze a resume, retrying temporary errors and falling back to other models."""
    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    models = [model] + [m for m in FALLBACK_MODELS if m != model]

    last_exc: Optional[Exception] = None
    for name in models:
        for attempt in range(MAX_RETRIES):
            try:
                return _generate(client, name, prompt)
            except errors.APIError as exc:
                last_exc = exc
                if exc.code == 404:  # model name not available: try the next one
                    break
                if exc.code not in RETRYABLE_CODES:
                    raise  # e.g. bad API key (400/403): retrying will not help
                if attempt < MAX_RETRIES - 1:
                    time.sleep(RETRY_BASE_DELAY * 2**attempt)
    assert last_exc is not None
    raise last_exc


def friendly_error(exc: Exception) -> str:
    code = getattr(exc, "code", None)
    if code in (429, 503):
        return (
            "Gemini is busy or you hit the rate limit. The app already retried and "
            "tried backup models. Please wait a minute and click Analyze again."
        )
    if code in (400, 401, 403):
        return "Gemini rejected the request. Check that your API key is valid."
    if code == 404:
        return "That Gemini model name was not found. Try 'gemini-flash-latest'."
    return f"Analysis failed: {exc}"


# ---------------------------------- UI -------------------------------------
def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: ResumeAnalysis) -> None:
    score = max(0, min(100, result.ats_score))
    st.divider()
    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric("ATS Score", f"{score}/100", score_label(score))
    with col2:
        st.progress(score / 100)
        st.write(result.summary)

    tab1, tab2, tab3, tab4 = st.tabs(
        ["Improvements", "Section scores", "Keywords", "Formatting & strengths"]
    )

    with tab1:
        order = {"high": 0, "medium": 1, "low": 2}
        for item in sorted(
            result.improvements, key=lambda i: order.get(i.priority.lower(), 3)
        ):
            icon = {"high": "🔴", "medium": "🟠", "low": "🟢"}.get(
                item.priority.lower(), "⚪"
            )
            with st.expander(f"{icon} {item.priority}: {item.issue}"):
                st.write(item.suggestion)

    with tab2:
        for sec in result.section_scores:
            s = max(0, min(100, sec.score))
            st.write(f"**{sec.name}** — {s}/100")
            st.progress(s / 100)
            st.caption(sec.feedback)

    with tab3:
        c1, c2 = st.columns(2)
        with c1:
            st.subheader("Found")
            st.write(", ".join(result.keywords_found) or "None detected")
        with c2:
            st.subheader("Missing")
            st.write(", ".join(result.keywords_missing) or "None")

    with tab4:
        st.subheader("Strengths")
        for s in result.strengths:
            st.markdown(f"- {s}")
        st.subheader("Formatting issues")
        if result.formatting_issues:
            for f in result.formatting_issues:
                st.markdown(f"- {f}")
        else:
            st.write("No major formatting issues found.")

    st.download_button(
        "Download report (JSON)",
        data=result.model_dump_json(indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and concrete improvements.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if not api_key:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.text_input(
            "Gemini model", value=os.environ.get("GEMINI_MODEL", DEFAULT_MODEL)
        )
        st.caption("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Max size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading resume..."):
                text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < 100:
            st.error(
                "Could not extract enough text. The file may be a scanned image. "
                "ATS systems also cannot read those, so export a text-based PDF or DOCX."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(
                    api_key, model.strip() or DEFAULT_MODEL, text, job_description
                )
        except Exception as exc:
            st.session_state.pop("result", None)
            st.error(friendly_error(exc))
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
